#!/usr/bin/env python3
"""
parse_calendar.py — Trauma On-Call Schedule Parser
====================================================
Converts a UofL Trauma-EGS Call Calendar Excel file (.xlsx) into
schedule.json for the Trauma On-Call web app.

Usage:
    python3 parse_calendar.py your_calendar_file.xlsx

Output:
    schedule.json  (written in the same directory as this script)

Requirements:
    pip install openpyxl

Expected Excel format:
    - One sheet per month (named "January", "February", etc.)
    - Each sheet has weekly blocks with:
        Row: dates (Sunday–Saturday in columns C–I)
        Row: Rounds - Tr1
        Row: Rounds - Tr2
        Row: Trauma Call  (or Trauma Call*)
        Row: Trauma Clinic   (Tuesday = Trauma Clinic, Wednesday = Burn Clinic)
        Row: Burn Clinic
        Row: Jewish Day
        Row: Jewish Night
        Row: Kindred
        Row: Norton ER
    - Blue font (RGB FF00B0F0) on Trauma Call = night call
"""

import sys
import json
import warnings
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is not installed.")
    print("Install it with:  pip install openpyxl")
    sys.exit(1)

warnings.filterwarnings('ignore')

# ── Column label → JSON key mapping ──────────────────────────────────────────
LABEL_MAP = {
    'rounds - tr1':  'tr1',
    'rounds - tr2':  'tr2',
    'trauma call*':  'trauma_call',
    'trauma call':   'trauma_call',
    'trauma clinic': 'trauma_clinic',   # Tue → trauma_clinic, Wed → burn_clinic
    'burn clinic':   'burn_clinic_raw', # usually blank; kept as fallback
    'jewish day':    'jewish_day',
    'jewish night':  'jewish_night',
    'kindred':       'kindred',
    'norton er':     'norton_er',
}

BLUE_RGB = 'FF00B0F0'   # blue font = trauma night call

MONTH_NAMES = [
    'January', 'February', 'March', 'April', 'May', 'June',
    'July', 'August', 'September', 'October', 'November', 'December'
]


def is_night_call(cell):
    """Return True if the cell's font is the blue night-call colour."""
    try:
        if cell.font and cell.font.color and cell.font.color.type == 'rgb':
            return cell.font.color.rgb.upper() == BLUE_RGB
    except Exception:
        pass
    return False


def parse_xlsx(xlsx_path: str) -> dict:
    """Parse the Excel file and return a dict keyed by 'YYYY-MM-DD'."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    schedule = {}

    for month_name in MONTH_NAMES:
        if month_name not in wb.sheetnames:
            print(f"  [skip] Sheet not found: {month_name}")
            continue

        ws = wb[month_name]
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=12))
        print(f"  Parsing {month_name}…", end=' ')
        week_count = 0

        i = 0
        while i < len(rows):
            row = rows[i]

            # Detect a date header row: ≥3 datetime values in columns C–I (indices 2–8)
            date_cells = {}
            for col_idx in range(2, 9):
                cell = row[col_idx] if col_idx < len(row) else None
                if cell and isinstance(cell.value, datetime):
                    date_cells[col_idx] = cell.value.date()

            if len(date_cells) >= 3:
                # Collect the data rows that follow until the next date block
                block_rows = {}
                seen_fields = set()   # stop if we loop back to a field we already collected

                j = i + 1
                while j < len(rows):
                    r = rows[j]

                    # Another date row → start of next weekly block
                    next_dates = [
                        r[k].value for k in range(2, 9)
                        if k < len(r) and isinstance(r[k].value, datetime)
                    ]
                    if len(next_dates) >= 3:
                        break

                    label_cell = r[1] if len(r) > 1 else None
                    label_val = str(label_cell.value or '').strip().lower() if label_cell else ''
                    field = LABEL_MAP.get(label_val)

                    if field:
                        if field in seen_fields:
                            break   # hit the next block's empty label rows — stop
                        seen_fields.add(field)
                        col_data = {}
                        for col_idx in range(2, 9):
                            if col_idx < len(r):
                                cell = r[col_idx]
                                val = str(cell.value or '').strip() if cell.value else ''
                                if val:
                                    col_data[col_idx] = (val, is_night_call(cell))
                        block_rows[field] = col_data

                    j += 1

                # Assign values to each date in this block
                for col_idx, dt in date_cells.items():
                    dt_str = dt.strftime('%Y-%m-%d')
                    if dt_str not in schedule:
                        schedule[dt_str] = {}
                    dow = dt.weekday()  # 0=Mon, 1=Tue, 2=Wed, …, 6=Sun

                    for field, col_data in block_rows.items():
                        if col_idx not in col_data:
                            continue
                        val, night = col_data[col_idx]

                        if field == 'trauma_call':
                            final_key = 'trauma_call_night' if night else 'trauma_call_day'
                        elif field == 'trauma_clinic':
                            # Wednesdays in the Trauma Clinic row are actually Burn Clinic
                            final_key = 'burn_clinic' if dow == 2 else 'trauma_clinic'
                        elif field == 'burn_clinic_raw':
                            final_key = 'burn_clinic'
                        else:
                            final_key = field

                        schedule[dt_str][final_key] = val

                week_count += 1
                i = j
            else:
                i += 1

        print(f"{week_count} week blocks parsed.")

    return schedule


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        print("Usage: python3 parse_calendar.py your_calendar_file.xlsx")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    if not Path(xlsx_path).exists():
        print(f"ERROR: File not found: {xlsx_path}")
        sys.exit(1)

    print(f"\nParsing: {xlsx_path}")
    schedule = parse_xlsx(xlsx_path)

    out_path = Path(__file__).parent / 'schedule.json'
    with open(out_path, 'w') as f:
        json.dump(schedule, f, separators=(',', ':'), sort_keys=True)

    print(f"\nDone! {len(schedule)} dates written to: {out_path}")
    print("Upload this file to GitHub to update the app.\n")


if __name__ == '__main__':
    main()
